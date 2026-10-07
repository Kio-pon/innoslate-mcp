"""The real browser path: headless Chromium, same-origin fetch, login form, UI clicks."""
from __future__ import annotations

import asyncio
import json

import pytest

from fake_innoslate import PASSWORD, USER
from innoslate_mcp import ui
from innoslate_mcp.api import InnoslateAPI
from innoslate_mcp.browser import BrowserSession, NotLoggedIn
from innoslate_mcp.plan import Planner
from innoslate_mcp.survey import Survey
from innoslate_mcp.urls import parse_project_url
from innoslate_mcp.workspace import workspace


@pytest.fixture
async def session():
    s = BrowserSession()
    yield s
    await s.close()


@pytest.fixture
def creds(monkeypatch):
    monkeypatch.setenv("INNOSLATE_USERNAME", USER)
    monkeypatch.setenv("INNOSLATE_PASSWORD", PASSWORD)


async def logged_in_api(session, fake) -> InnoslateAPI:
    p = parse_project_url(fake.project_url)
    r = await session.login(p, 30)
    assert r["logged_in"], r
    return InnoslateAPI(session, p)


@pytest.mark.parametrize("style", ["single", "two-step"])
async def test_login_with_stored_credentials(session, fake, creds, style):
    fake.store.login_style = style
    p = parse_project_url(fake.project_url)
    r = await session.login(p, 30)
    assert r == {"logged_in": True, "status": 200, "how": "stored credentials typed into the login form"}
    again = await session.login(p, 5)
    assert again["how"] == "already logged in"
    assert "headless" in session.mode


async def test_wrong_password_times_out_cleanly(session, fake, monkeypatch):
    monkeypatch.setenv("INNOSLATE_USERNAME", USER)
    monkeypatch.setenv("INNOSLATE_PASSWORD", "wrong")
    r = await session.login(parse_project_url(fake.project_url), 6)
    assert r["logged_in"] is False and "did not work" in r["how"]


async def test_user_logs_in_by_hand_while_server_waits(session, fake):
    p = parse_project_url(fake.project_url)

    async def human():
        await asyncio.sleep(2)
        page = await session.new_page()
        await page.goto(fake.base + "/login?next=/")
        await page.fill("input[name=username]", USER)
        await page.fill("input[name=password]", PASSWORD)
        await page.click("button[type=submit]")

    r, _ = await asyncio.gather(session.login(p, 30), human())
    assert r["logged_in"] and r["how"] == "user logged in in the browser window"


@pytest.mark.parametrize("typed", [False, True])
async def test_sso_login_on_another_origin_opens_no_extra_tabs(session, fake, monkeypatch, typed):
    fake.store.login_style = "sso"
    if typed:
        monkeypatch.setenv("INNOSLATE_USERNAME", USER)
        monkeypatch.setenv("INNOSLATE_PASSWORD", PASSWORD)
    p = parse_project_url(fake.project_url)

    async def human():
        await asyncio.sleep(5)                               # several polls happen while the user is on the SSO page
        ctx = await session.context()
        pages = [x for x in ctx.pages if "localhost" in x.url]
        if typed:
            return len(ctx.pages)
        assert len(pages) == 1, [x.url for x in ctx.pages]
        await pages[0].fill("input[name=username]", USER)
        await pages[0].fill("input[name=password]", PASSWORD)
        await pages[0].click("button[type=submit]")
        return len(ctx.pages)

    r, tabs = await asyncio.gather(session.login(p, 30), human())
    assert r["logged_in"], r
    assert tabs == 1 and len((await session.context()).pages) == 1
    api = InnoslateAPI(session, p)
    assert (await api.get(2101))["number"] == "UN.1.1"


async def test_expired_login_raises_clear_error(session, fake, creds):
    api = await logged_in_api(session, fake)
    assert (await api.get(2101))["number"] == "UN.1.1"
    fake.store.sessions.clear()
    with pytest.raises(NotLoggedIn, match="innoslate_login"):
        await api.get(2101)
    ok, status = await session.is_logged_in(api.p)
    assert not ok and status == 401


async def test_survey_writes_map_and_nomenclature(session, fake, creds, tmp_path):
    import docx
    g = tmp_path / "guides"
    g.mkdir()
    d = docx.Document()
    d.add_paragraph('Number the document SR.1 and each requirement SR.1.n. User needs go in UN.1 as UN.n.n.')
    d.add_paragraph('Use the Requirements Document label. Create ACT.1 "Top-level action diagram". Traced From links.')
    d.save(g / "Deliverable 3 guide.docx")
    api = await logged_in_api(session, fake)
    out = await Survey(api, workspace(api.p)).run(str(g))
    assert out["entities"] == len(fake.store.E)
    clashes = [(c["number"], c["existing_id"]) for c in out["numbering_clashes"]]
    assert sorted(clashes) == [("SR.1", 1002), ("UN.1", 1001), ("UN.1", 2001)]
    pats = {p["pattern"] for p in out["guide_patterns"]}
    assert {"SR.1", "SR.1.n", "UN.n.n", "ACT.1"} <= pats
    assert any(p["problem"] == "Relations recorded on one side only" for p in out["problems"])
    assert out["comments"] == {"total": 2, "unresolved": 1}
    assert out["storage_conventions"]["item_to_document_relations"].get("sourced by", 0) >= 6
    m = (workspace(api.p) / "PROJECT_MAP.md").read_text()
    assert "### `UN.1` MP. User Needs Document (id 2001)" in m
    assert "  - `UN.1.1.1` **Switch areas**" in m                       # nested item shown under its parent
    assert "UN.1.3 is not atomic" in m and "instructor@uni.example" not in m
    assert "Numbers used more than once (" in m and "UN.1.1" in m
    n = (workspace(api.p) / "NOMENCLATURE.md").read_text()
    assert "`SR.1` is already 'System Requirements (senior team)'" in n
    assert "Labels: " in n and "Requirements Document" in n
    assert "Relations: " in n and "traced from" in n
    snap = json.loads((workspace(api.p) / "snapshot.json").read_text())
    art = next(e for e in snap["entities"] if e["n"] == "ART.1")
    assert art["attrs"]["File"].startswith("FILE:")


async def test_write_through_the_browser(session, fake, creds):
    api = await logged_in_api(session, fake)
    plan = {"entities": [
        {"key": "sr", "class": "Artifact", "number": "SR.3", "name": "Browser test doc", "labels": ["Requirements Document"]},
        {"key": "r", "class": "Requirement", "number": "SR.3.1", "name": "Alerts", "document": "sr",
         "description": "The system shall list every open alert.", "relations": {"traced from": [{"number": "UN.1.2", "document": 2001}]}},
    ]}
    pl = Planner(api, api.p.pid)
    res = await pl.commit(await pl.compile(plan))
    assert res["verify"]["ok"], res
    assert len(fake.store.posts) == 1 and len(fake.store.posts[0]) == 3   # doc, requirement, UN.1.2 (inverse side)


async def test_upload_through_the_browser(session, fake, creds, tmp_path):
    api = await logged_in_api(session, fake)
    f = tmp_path / "Visit notes.pdf"
    payload = bytes(range(256)) * 50
    f.write_bytes(payload)
    key = await api.upload_file(f)
    assert fake.store.files[key]["data"] == payload
    assert fake.store.files[key]["name"] == "Visit notes.pdf" and fake.store.files[key]["mime"] == "application/pdf"


async def test_report_download_and_wrong_name(session, fake, creds, tmp_path):
    api = await logged_in_api(session, fake)
    r = await ui.download_report(session, api, 2001, "RTM Output (XLSX)", tmp_path / "out")
    assert r["ok"] and r["file"].endswith("RTM Output.xlsx")
    assert (tmp_path / "out" / "RTM Output.xlsx").read_bytes() == b"PK\x03\x04fake-report-2001"
    r2 = await ui.download_report(session, api, 2001, "RTM Output (XLSX)", tmp_path / "out")
    assert r2["ok"] and r2["file"] != r["file"]                         # never overwrites an earlier download
    bad = await ui.download_report(session, api, 2001, "RTM Output", tmp_path / "out")
    assert not bad["ok"] and "RTM Output (XLSX)" in bad["available_reports"] and bad["did_you_mean"]


async def test_quality_check(session, fake, creds):
    api = await logged_in_api(session, fake)
    plan = {"entities": [
        {"key": "sr", "class": "Artifact", "number": "SR.4", "name": "QC doc", "labels": ["Requirements Document"]},
        {"key": "a", "class": "Requirement", "number": "SR.4.1", "name": "Good", "document": "sr",
         "description": "The system shall list every open alert."},
        {"key": "b", "class": "Requirement", "number": "SR.4.2", "name": "Bad", "document": "sr",
         "description": "The system should be fast."},
    ]}
    pl = Planner(api, api.p.pid)
    c = await pl.compile(plan)
    await pl.commit(c)
    doc_id = next(e["id"] for e in fake.store.E.values() if e["number"] == "SR.4")
    r = await ui.quality_check(session, api, doc_id, 60)
    assert r["ok"] and r["scored"] == 2 and fake.store.qc_runs == 1
    assert r["lowest"][0]["number"] == "SR.4.2" and r["lowest"][0]["score"] == 25
    assert {f["check"] for f in r["lowest"][0]["failed_checks"]} == {"Unambiguous", "Singular", "Verifiable"}
    assert "Quality check finished for 2" in r["dialog_text"]
    none = await ui.quality_check(session, api, 2001, 10)          # a needs document has no requirements
    assert not none["ok"] and "Nothing to check" in none["error"]


async def test_closed_window_is_reopened(session, fake, creds):
    api = await logged_in_api(session, fake)
    for p in list((await session.context()).pages):
        await p.close()
    assert (await api.get(2101))["number"] == "UN.1.1"


async def test_missing_chromium_triggers_one_install(monkeypatch, tmp_path):
    """With no browser on disk the session calls the installer once, then reports a clear error."""
    import innoslate_mcp.browser as b
    calls = []

    def fake_install():
        calls.append(1)
        raise b.InnoslateError("Could not install Chromium. Run this once by hand: python -m playwright install chromium")

    monkeypatch.setenv("PLAYWRIGHT_BROWSERS_PATH", str(tmp_path / "no-browsers"))
    monkeypatch.setattr(b, "install_chromium", fake_install)
    s = b.BrowserSession()
    try:
        with pytest.raises(b.InnoslateError, match="playwright install chromium"):
            await s.context()
    finally:
        await s.close()
    assert calls == [1]


async def test_find_number_falls_back_to_snapshot(fake_nonum, monkeypatch):
    """Over MCP-level code: search that ignores numbers still finds UN.1.1 after initialize."""
    from innoslate_mcp import server
    monkeypatch.setenv("INNOSLATE_USERNAME", USER)
    monkeypatch.setenv("INNOSLATE_PASSWORD", PASSWORD)
    server.STATE.__init__()
    try:
        await server.innoslate_login(fake_nonum.project_url, 30)
        assert (await server.innoslate_find(number="UN.1.1"))["total"] == 0      # no snapshot yet
        await server.innoslate_initialize()
        assert (await server.innoslate_find(number="UN.1.1"))["total"] == 2
        assert (await server.innoslate_get("UN.1.2"))["matches"] == 2
    finally:
        await server.STATE.session.close()
