"""Pure functions: URL parsing, schema resolution, lint, local document reading."""
from __future__ import annotations

import zipfile

import pytest

from fake_innoslate import CLASSES, LABELS, PROPERTIES, RELATIONS
from innoslate_mcp import config, docs, lint
from innoslate_mcp.api import strip_html, to_html
from innoslate_mcp.schema import Schema
from innoslate_mcp.urls import parse_project_url

SCHEMA = Schema({"classes": CLASSES, "relations": RELATIONS, "labels": LABELS, "properties": PROPERTIES})


# ---------------------------------------------------------------------- urls
@pytest.mark.parametrize("url,base,org,pid", [
    ("https://cloud.innoslate.com/myorg/p/1234/dashboard", "https://cloud.innoslate.com", "myorg", 1234),
    ("https://cloud.innoslate.com/myorg/p/1234", "https://cloud.innoslate.com", "myorg", 1234),
    ("https://cloud.innoslate.com/myorg/p/1234/documents/567890?x=1", "https://cloud.innoslate.com", "myorg", 1234),
    ("  https://in.example.org/innoslate/acme/p/12/database#top ", "https://in.example.org/innoslate", "acme", 12),
    ("http://127.0.0.1:8765/acme/p/42/dashboard", "http://127.0.0.1:8765", "acme", 42),
])
def test_parse_project_url(url, base, org, pid):
    p = parse_project_url(url)
    assert (p.base, p.org, p.pid) == (base, org, pid)
    assert p.papi == f"{base}/api/v4/o/{org}/p/{pid}"
    assert p.document_url(7) == f"{base}/{org}/p/{pid}/documents/7"
    assert p.origin.startswith("http")


@pytest.mark.parametrize("bad", ["", "cloud.innoslate.com/myorg/p/1", "https://cloud.innoslate.com/myorg/dashboard",
                                 "https://cloud.innoslate.com/myorg/p/abc"])
def test_parse_project_url_rejects(bad):
    with pytest.raises(ValueError, match="Innoslate project URL"):
        parse_project_url(bad)


# -------------------------------------------------------------------- schema
def test_schema_names_case_and_space_insensitive():
    assert SCHEMA.class_id("  requirement ") == 22
    assert SCHEMA.class_id("Input/Output") == 14
    assert SCHEMA.relation_id("Traced  From") == 81
    assert SCHEMA.inverse_id(81) == 82
    assert SCHEMA.inverse_id(78) == 77
    assert SCHEMA.label_id("functional requirement") == 103


def test_schema_unknown_name_suggests():
    with pytest.raises(LookupError, match="Closest: traced from"):
        SCHEMA.relation_id("traced frm")
    with pytest.raises(LookupError, match="Requirement"):
        SCHEMA.class_id("Requirment")


def test_schema_property_through_parent_class():
    # Test Case (24) inherits from Action (1), which owns Duration
    assert SCHEMA.class_ancestry(24) == [24, 1]
    assert SCHEMA.property(24, "duration")["id"] == 40
    with pytest.raises(LookupError):
        SCHEMA.property(2, "Rationale")


def test_schema_label_fit_and_summary():
    assert SCHEMA.label_fits(22, 2)
    assert not SCHEMA.label_fits(103, 2)
    s = SCHEMA.summary("all")
    assert "traced from <-> traced to" in s["relations"]
    props = SCHEMA.summary("properties", "Requirement")["properties"]
    qs = next(p for p in props if p["name"] == "Quality Score")
    assert qs["writable"] is False


def test_schema_accepts_dict_shaped_lists():
    s = Schema({"classes": {str(c["id"]): c for c in CLASSES}, "relations": RELATIONS})
    assert s.class_id("Artifact") == 2


# ---------------------------------------------------------------------- html
def test_html_roundtrip():
    h = to_html("First <line> & more\n\nSecond")
    assert h == "<p>First &lt;line&gt; &amp; more</p><p>Second</p>"
    assert strip_html(h) == "First <line> & more\nSecond"
    assert strip_html("<table><tr><td>a</td><td>b</td></tr></table>") == "a | b |"


# ---------------------------------------------------------------------- lint
def test_lint_catches_common_problems():
    r = lint.lint_items([
        {"number": "SR.1.1", "name": "Good", "text": "The system shall display the alert list."},
        {"number": "SR.1.2", "name": "Bad", "text": "The system should be fast and user-friendly, etc."},
        {"number": "SR.1.3", "name": "Two", "text": "The system shall log events. The system shall export them."},
        {"number": "SR.1.4", "name": "Num", "text": "The system shall refresh within 5 seconds."},
        {"number": "SR.1.4", "name": "Dup", "text": "The system shall display the alert list."},
    ])
    by = {i["number"] + i["name"]: [f["message"] for f in i["findings"]] for i in r["items"]}
    assert r["clean"] == ["SR.1.1"]
    bad = " ".join(by["SR.1.2Bad"])
    assert "should" in bad and "fast" in bad and "etc" in bad and "no 'shall'" in bad
    assert any("more than one sentence" in m for m in by["SR.1.3Two"])
    assert any("5 seconds" in m for m in by["SR.1.4Num"])
    assert any("same text as SR.1.1" in m for m in by["SR.1.4Dup"])
    assert any("number used more than once" in m for m in by["SR.1.4Dup"])
    assert r["errors"] >= 5


def test_lint_extra_banned():
    r = lint.lint_items([{"number": "X", "text": "The system shall use facial recognition."}],
                        extra_banned={r"\bfacial\b": "no worker identification"})
    assert any("no worker identification" in f["message"] for f in r["items"][0]["findings"])


# ---------------------------------------------------------------------- docs
def _tiny_pdf(text: str) -> bytes:
    """A valid one-page PDF with real text, built by hand so no PDF writer is needed."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def test_read_pdf(tmp_path):
    p = tmp_path / "guide.pdf"
    p.write_bytes(_tiny_pdf("Number stakeholder requirements SR.1.n"))
    t = docs.extract_text(p)
    assert "--- page 1 ---" in t and "SR.1.n" in t


def test_read_docx_keeps_headings_and_tables(tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("User Needs", 1)
    d.add_paragraph("Supervisors need alerts.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "ID", "Need"
    t.cell(1, 0).text, t.cell(1, 1).text = "UN.1.1", "Live view"
    d.add_paragraph("After the table.")
    p = tmp_path / "needs.docx"
    d.save(p)
    text = docs.extract_text(p)
    assert text.index("# User Needs") < text.index("| UN.1.1 | Live view |") < text.index("After the table.")


def test_read_xlsx_csv_json_vtt_pptx(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "RTM"
    wb.active.append(["SR", "UN"])
    wb.active.append(["SR.1.1", "UN.1.1"])
    wb.save(tmp_path / "m.xlsx")
    assert "| SR.1.1 | UN.1.1 |" in docs.extract_text(tmp_path / "m.xlsx")
    (tmp_path / "a.csv").write_text("a,b\n1,2\n")
    assert docs.extract_text(tmp_path / "a.csv").splitlines()[1] == "| 1 | 2 |"
    (tmp_path / "a.json").write_text('{"k":[1]}')
    assert '"k"' in docs.extract_text(tmp_path / "a.json")
    (tmp_path / "m.vtt").write_text("WEBVTT\n\n00:01:05.000 --> 00:01:09.000\nThe camera in area two failed.\n")
    assert docs.extract_text(tmp_path / "m.vtt") == "[01:05] The camera in area two failed."
    with zipfile.ZipFile(tmp_path / "s.pptx", "w") as z:
        z.writestr("ppt/slides/slide1.xml", "<p:sld><a:p><a:r><a:t>Scope &amp; goals</a:t></a:r></a:p></p:sld>")
        z.writestr("ppt/slides/slide2.xml", "<p:sld><a:p><a:r><a:t>Second</a:t></a:r></a:p></p:sld>")
    assert docs.extract_text(tmp_path / "s.pptx") == "--- slide 1 ---\nScope & goals\n\n--- slide 2 ---\nSecond"


def test_read_page_paging_and_errors(tmp_path):
    p = tmp_path / "long.txt"
    p.write_text("x" * 2500)
    a = docs.read_page(p, 0, 1000)
    assert a["returned_chars"] == 1000 and a["next_offset"] == 1000
    c = docs.read_page(p, 2000, 1000)
    assert c["next_offset"] is None and c["returned_chars"] == 500
    with pytest.raises(FileNotFoundError):
        docs.extract_text(tmp_path / "nope.pdf")
    (tmp_path / "v.mp4").write_bytes(b"\x00\x00\x00 ftypmp42" + b"\x00" * 100)
    with pytest.raises(docs.UnsupportedFile, match="transcript"):
        docs.extract_text(tmp_path / "v.mp4")


def test_list_files_skips_hidden(tmp_path):
    (tmp_path / "a.pdf").write_bytes(b"%PDF")
    (tmp_path / ".secret").write_text("x")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "b.docx").write_bytes(b"PK")
    from pathlib import Path
    names = [Path(f["path"]).name for f in docs.list_files(tmp_path)]
    assert names == ["a.pdf", "b.docx"]


# -------------------------------------------------------------------- config
def test_credentials_from_env_file_and_repr_hides_password(isolated_home, monkeypatch):
    isolated_home.mkdir(parents=True, exist_ok=True)
    (isolated_home / ".env").write_text("# mine\nexport INNOSLATE_USERNAME=me@uni.edu\nINNOSLATE_PASSWORD='p@ss word'\n")
    c = config.credentials()
    assert c.username == "me@uni.edu" and c.password == "p@ss word"
    assert "p@ss" not in repr(c) and "p@ss" not in str(c)
    monkeypatch.setenv("INNOSLATE_USERNAME", "override")
    assert config.credentials().username == "override"


@pytest.mark.skipif(__import__("os").name != "posix", reason="file modes are POSIX only")
def test_env_file_permission_warning(isolated_home):
    import os
    isolated_home.mkdir(parents=True, exist_ok=True)
    f = isolated_home / ".env"
    f.write_text("INNOSLATE_PASSWORD=x\n")
    os.chmod(f, 0o644)
    assert "chmod 600" in (config.env_file_warning() or "")
    os.chmod(f, 0o600)
    assert config.env_file_warning() is None
