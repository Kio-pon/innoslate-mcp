"""Read the user's local documents as plain text: PDF, DOCX, XLSX, PPTX, transcripts, text.

Paging (offset, max_chars) keeps a long guide from flooding the AI's context.
"""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from pathlib import Path

TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".log", ".tex", ".html", ".htm", ".xml", ".yaml", ".yml"}
SUPPORTED = TEXT_EXT | {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".csv", ".tsv", ".json", ".vtt", ".srt"}


class UnsupportedFile(ValueError):
    pass


def _pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise UnsupportedFile("pypdf is not installed. Run: pip install pypdf") from e
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:
            raise UnsupportedFile(f"{path.name} is password protected.") from e
    pages = []
    for i, p in enumerate(reader.pages, 1):
        try:
            t = p.extract_text() or ""
        except Exception:
            t = ""
        pages.append(f"--- page {i} ---\n{t.strip()}")
    text = "\n\n".join(pages)
    if len(re.sub(r"--- page \d+ ---|\s", "", text)) < 20 * max(1, len(reader.pages)):
        text += ("\n\n[Almost no text could be extracted. This PDF is probably scanned images. "
                 "Ask the user for a text version or an OCR'd copy.]")
    return text


def _docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as e:  # pragma: no cover
        raise UnsupportedFile("python-docx is not installed. Run: pip install python-docx") from e
    d = docx.Document(str(path))
    out: list[str] = []
    # Walk the body in order so tables stay next to the paragraphs around them.
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.lower().startswith("heading"):
                lvl = re.sub(r"\D", "", style) or "1"
                out.append("#" * min(int(lvl), 6) + " " + t)
            elif "list" in style.lower():
                out.append("- " + t)
            else:
                out.append(t)
        elif tag == "tbl":
            tbl = Table(child, d)
            rows = []
            for r in tbl.rows:
                cells = []
                for c in r.cells:
                    cells.append(" ".join(c.text.split()))
                rows.append("| " + " | ".join(cells) + " |")
            if rows:
                out.append("\n".join(rows))
    for s in d.sections:
        for part, name in ((s.header, "header"), (s.footer, "footer")):
            try:
                t = " ".join(p.text for p in part.paragraphs).strip()
            except Exception:
                t = ""
            if t:
                out.append(f"[{name}] {t}")
                break
    return "\n\n".join(out)


def _xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover
        raise UnsupportedFile("openpyxl is not installed. Run: pip install openpyxl") from e
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if row is None or all(v in (None, "") for v in row):
                continue
            out.append("| " + " | ".join("" if v is None else str(v).replace("\n", " ") for v in row) + " |")
    wb.close()
    return "\n".join(out)


def _pptx(path: Path) -> str:
    """Slide text straight from the XML, so python-pptx is not needed."""
    out = []
    with zipfile.ZipFile(path) as z:
        slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                        key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)))
        for i, n in enumerate(slides, 1):
            xml = z.read(n).decode("utf-8", "replace")
            paras = []
            for para in re.findall(r"<a:p>(.*?)</a:p>", xml, flags=re.S):
                t = "".join(re.findall(r"<a:t>(.*?)</a:t>", para, flags=re.S))
                if t.strip():
                    paras.append(_unxml(t))
            out.append(f"--- slide {i} ---\n" + "\n".join(paras))
    return "\n\n".join(out)


def _unxml(t: str) -> str:
    import html
    return html.unescape(t)


def _subtitles(text: str) -> str:
    """VTT/SRT to 'mm:ss text' lines."""
    out = []
    stamp = None
    for line in text.splitlines():
        line = line.strip()
        m = re.match(r"(\d{1,2}:)?(\d{2}):(\d{2})[.,]\d{3}\s*-->", line)
        if m:
            h = int((m.group(1) or "0:")[:-1])
            stamp = f"{h * 60 + int(m.group(2)):02d}:{m.group(3)}"
            continue
        if not line or line.isdigit() or line.upper().startswith("WEBVTT") or line.startswith("NOTE"):
            continue
        line = re.sub(r"<[^>]+>", "", line)
        out.append(f"[{stamp}] {line}" if stamp else line)
        stamp = None
    return "\n".join(out)


def extract_text(path: str | Path) -> str:
    p = Path(path).expanduser()
    if not p.exists():
        raise FileNotFoundError(f"No such file: {p}")
    if p.is_dir():
        raise UnsupportedFile(f"{p} is a folder. Call list_local_files on it.")
    ext = p.suffix.lower()
    if ext == ".pdf":
        return _pdf(p)
    if ext == ".docx":
        return _docx(p)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx(p)
    if ext == ".pptx":
        return _pptx(p)
    raw = p.read_bytes()
    text = raw.decode("utf-8-sig", errors="replace")
    if ext in (".vtt", ".srt"):
        return _subtitles(text)
    if ext == ".json":
        try:
            return json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except ValueError:
            return text
    if ext in (".csv", ".tsv"):
        rows = csv.reader(io.StringIO(text), delimiter="\t" if ext == ".tsv" else ",")
        return "\n".join("| " + " | ".join(r) + " |" for r in rows)
    if ext in TEXT_EXT or b"\x00" not in raw[:4096]:
        if ext in (".html", ".htm"):
            from .api import strip_html
            return strip_html(text)
        return text
    raise UnsupportedFile(f"Cannot read {p.suffix or 'this'} files. Supported: {', '.join(sorted(SUPPORTED))}. "
                          "Audio and video need a transcript first.")


def read_page(path: str | Path, offset: int = 0, max_chars: int = 30000) -> dict:
    text = extract_text(path)
    offset = max(0, int(offset))
    max_chars = max(1000, min(int(max_chars), 200000))
    chunk = text[offset:offset + max_chars]
    end = offset + len(chunk)
    return {
        "path": str(Path(path).expanduser().resolve()),
        "total_chars": len(text),
        "offset": offset,
        "returned_chars": len(chunk),
        "next_offset": end if end < len(text) else None,
        "text": chunk,
    }


def list_files(folder: str | Path, recursive: bool = True) -> list[dict]:
    root = Path(folder).expanduser()
    if not root.is_dir():
        raise FileNotFoundError(f"No such folder: {root}")
    it = root.rglob("*") if recursive else root.glob("*")
    out = []
    for f in sorted(it):
        if f.is_file() and not any(part.startswith(".") for part in f.relative_to(root).parts):
            out.append({"path": str(f), "size_kb": round(f.stat().st_size / 1024, 1),
                        "readable": f.suffix.lower() in SUPPORTED})
        if len(out) >= 500:
            break
    return out
